#!/usr/bin/env python3
"""
Convert the Maharashtra State CET Cell "Provisional Seat Matrix for MBBS/BDS Courses"
(college-wise vacancy) PDF into an Excel workbook that keeps the PDF's layout.

Usage:
    pip install pdfplumber openpyxl
    python seat_matrix_pdf_to_excel.py input.pdf [output.xlsx]            # layout like the PDF
    python seat_matrix_pdf_to_excel.py input.pdf out.xlsx --long          # one row per seat entry
        (columns: Sr, Code, College, Aided, Min, GOI, <GEN./WOM./PWD..>, Category, Seats)
    options for --long:  --keep-zeros   keep cells printed as 0 (default: only seats > 0)
                         --section-col  add a 'Section' column (Govt/Pvt/Pvt Univ + course)
                         --csv          also write a tab-separated .txt/.csv next to the xlsx

Output workbook:
    * 'Checks' sheet (first) - live Excel formulas that validate the conversion
    * one sheet per section (Govt MBBS, Pvt MBBS, Pvt Univ MBBS, Govt BDS, ...)
    * a validation report is also printed in the terminal

How it works: the PDF is a fixed-width text report, so every number is mapped to a
column by the x-position of its right edge (GRID below). If CET changes the report
layout, adjust GRID (run with --debug-grid to see the x-positions found in your PDF).
"""
import re
import sys
import argparse
import collections

import pdfplumber
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

# ----------------------------------------------------------------------------
# 1. PDF PARSING
# ----------------------------------------------------------------------------
# right-edge x-position (PDF points) of each numeric column
GRID = {
    49.7: 'Sr', 72.2: 'Code', 288.2: 'Cap', 306.2: 'AIQ', 355.7: 'GOI',
    396.2: 'SC', 414.2: 'ST', 432.2: 'VJA', 450.2: 'NTB', 468.2: 'NTC', 486.2: 'NTD',
    504.2: 'OBC', 526.6: 'CTot', 544.6: 'SEB', 562.6: 'EWS', 585.1: 'OPEN',
    612.2: 'StateTot', 634.7: 'StateTotCol',
    657.2: 'D1', 670.7: 'D2', 684.2: 'D3', 697.7: 'MK',
    715.7: 'IQ',          # IQ  (or NRI in "university" sections)
    733.7: 'AIQ2',        # AIQ (university sections only)
    751.7: 'OthTot', 769.7: 'OthTotCol', 792.2: 'Bal', 814.7: 'BalCol',
}
LABELS = {'GEN.', 'WOM.', 'PWD.', 'ORP.', 'ORPC'}   # 'GEN HA' / 'WOMEN HA' handled separately
BODY_TOP, BODY_BOTTOM = 88, 548                      # ignore page header / footer legends


def snap(x, tol=3.0):
    k = min(GRID, key=lambda g: abs(g - x))
    return GRID[k] if abs(k - x) <= tol else None


def numeric(word):
    """Return (value, right_edge_x) if the word is an integer (optionally glued to '|')."""
    t = word['text']
    if t.endswith('|') and t[:-1].isdigit():
        return int(t[:-1]), word['x1'] - 4.5
    if t.isdigit():
        return int(t), word['x1']
    return None


def page_lines(page):
    words = [w for w in page.extract_words() if BODY_TOP < w['top'] < BODY_BOTTOM]
    words.sort(key=lambda w: (w['top'], w['x0']))
    lines, cur, top = [], [], None
    for w in words:
        if top is None or abs(w['top'] - top) > 2.5:
            if cur:
                lines.append(cur)
            cur, top = [w], w['top']
        else:
            cur.append(w)
    if cur:
        lines.append(cur)
    return lines


def parse_pdf(path, debug_grid=False):
    sections, sec, col = [], None, None
    with pdfplumber.open(path) as pdf:
        if debug_grid:
            c = collections.Counter()
            for p in pdf.pages:
                for w in p.extract_words():
                    n = numeric(w)
                    if n and w['top'] > BODY_TOP:
                        c[round(n[1], 1)] += 1
            print('x-positions of numbers found:', sorted(c.items()))
        for pi, page in enumerate(pdf.pages):
            for ln in page_lines(page):
                toks = [w['text'] for w in ln]
                txt = ' '.join(toks)
                if all(set(t) <= set('-=') for t in toks):            # dashed separators
                    continue
                if 'COLLEGES:' in toks:                               # section title
                    name = ' '.join(t for t in toks if t not in ('Status', ':', 'Original'))
                    if sections and sections[-1]['name'] == name:
                        sec = sections[-1]
                    else:
                        sec = {'name': name, 'colleges': [], 'total': None, 'req': None, 'req_text': None}
                        sections.append(sec)
                    col = None
                    continue
                if toks[0] == 'Req.':                                 # "Req. Of:5065 SC ST ..."
                    sec['req_text'] = ('Req. Of: ' + toks[2]) if toks[1] == 'Of:' else ('Req. ' + toks[1])
                    sec['req'] = int(re.search(r'\d+', sec['req_text']).group())
                    continue
                if 'Cap.' in toks and 'SC' in toks and sec is not None:   # header line 2: remember this PDF's own labels
                    i0 = toks.index('SC')
                    sec['catnames'] = dict(zip(['SC', 'ST', 'VJA', 'NTB', 'NTC', 'NTD', 'OBC'], toks[i0:i0 + 7]))
                    # layout detection: where does the last "Tot" of the header sit?
                    #   x1 ~ 751.7 -> wide layout (Round 3 PDF), x1 ~ 733.7 -> compact layout (Group A PDF)
                    tot_x = [w['x1'] for w in ln if w['text'].startswith('Tot')][-1]
                    sec['layout'] = 'compact' if tot_x < 742 else 'wide'
                    continue
                if toks[0] in ('Sr', 'Govt/', 'Aided', 'Min', 'PVT.') or 'Cap.' in toks:
                    continue                                          # column header lines

                # category label on this line
                label = next((w['text'] for w in ln if 340 <= w['x0'] <= 380 and w['text'] in LABELS), None)
                if label is None:
                    ha = [w for w in ln if 340 <= w['x0'] <= 372 and w['text'] in ('GEN', 'WOMEN')]
                    if ha:
                        label = ha[0]['text'] + ' HA'
                is_total = any(w['text'] == '*' for w in ln)
                name_parts = [w['text'] for w in ln
                              if 80 <= w['x0'] <= 300 and w['x1'] <= 246 and w['text'] != '*' and numeric(w) is None]

                nums = []
                for w in ln:
                    n = numeric(w)
                    if n is None or (340 <= w['x0'] <= 380 and w['text'] in LABELS):
                        continue
                    c = snap(n[1])
                    if sec is not None and sec.get('layout') == 'compact':
                        c = {'AIQ2': 'OthTot', 'OthTot': 'OthTotCol'}.get(c, c)
                    nums.append((c, n[0], n[1]))
                has_sr = any(c == 'Sr' for c, _, _ in nums)
                has_code = any(c == 'Code' for c, _, _ in nums)

                if has_sr and has_code and not is_total:              # first line of a college
                    col = {
                        'sr': next(v for c, v, _ in nums if c == 'Sr'),
                        'code': next(v for c, v, _ in nums if c == 'Code'),
                        'name': ' '.join(name_parts),
                        'gov': next((w['text'] for w in ln if 245 <= w['x0'] <= 256 and w['text'] in ('Y', 'N')), None),
                        'min': ' '.join(w['text'] for w in ln if 308 <= w['x0'] <= 340 and w['x1'] <= 340 and numeric(w) is None),
                        'rows': [],
                    }
                    sec['colleges'].append(col)
                elif is_total:
                    col = {'total': True, 'rows': []}
                    sec['total'] = col
                elif label is None:                                   # wrapped college name (3rd line etc.)
                    if col is not None and name_parts and not col.get('total'):
                        col['name'] += ' ' + ' '.join(name_parts)
                    continue
                elif col is not None and name_parts and not col.get('total'):   # name wraps onto data line
                    col['name'] += ' ' + ' '.join(name_parts)

                if label is None:
                    print(f'WARNING page {pi + 1}: line without category label: {txt}')
                    continue
                vals = {}
                for c, v, x in nums:
                    if c in ('Sr', 'Code'):
                        continue
                    if c is None:
                        print(f'WARNING page {pi + 1}: number {v} at x={x:.1f} matches no column: {txt}')
                        continue
                    vals[c] = v
                col['rows'].append({'label': label, 'vals': vals})
    return sections


# ----------------------------------------------------------------------------
# 2. VALIDATION IN PYTHON (printed in terminal)
# ----------------------------------------------------------------------------
def validate(sections):
    g = lambda v, k: v.get(k, 0)
    cat = ['SC', 'ST', 'VJA', 'NTB', 'NTC', 'NTD', 'OBC']
    oth = ['D1', 'D2', 'D3', 'MK', 'IQ', 'AIQ2']
    problems = 0
    for s in sections:
        sp = 0
        for c in s['colleges']:
            ss = so = sb = 0
            for r in c['rows']:
                v = r['vals']
                if sum(g(v, k) for k in cat) != g(v, 'CTot'):
                    print('  Tot. mismatch', c['code'], r['label']); sp += 1
                if g(v, 'CTot') + g(v, 'SEB') + g(v, 'EWS') + g(v, 'OPEN') != g(v, 'StateTot'):
                    print('  State Tot mismatch', c['code'], r['label']); sp += 1
                if sum(g(v, k) for k in oth) != g(v, 'OthTot'):
                    print('  Other Tot mismatch', c['code'], r['label']); sp += 1
                if g(v, 'StateTot') + g(v, 'OthTot') != g(v, 'Bal'):
                    print('  Bal mismatch', c['code'], r['label']); sp += 1
                ss += g(v, 'StateTot'); so += g(v, 'OthTot'); sb += g(v, 'Bal')
            f = c['rows'][0]['vals']
            for key, tot, nm in (('StateTotCol', ss, 'State'), ('OthTotCol', so, 'Other'), ('BalCol', sb, 'Bal')):
                if f.get(key) != tot:
                    print(f'  College {nm} total mismatch', c['code'], f.get(key), tot); sp += 1
        for r in s['total']['rows']:
            for k, v in r['vals'].items():
                if k in ('Cap', 'AIQ', 'GOI'):
                    comp = sum(g(c['rows'][0]['vals'], k) for c in s['colleges'])
                    if r is not s['total']['rows'][0]:
                        continue
                else:
                    comp = sum(g(rr['vals'], k) for c in s['colleges'] for rr in c['rows'] if rr['label'] == r['label'])
                if comp != v:
                    print(f"  Total row mismatch {r['label']} {k}: printed {v}, computed {comp}"); sp += 1
        t = s['total']['rows'][0]['vals']
        if s['req'] is not None and s['req'] != t.get('Cap', 0) - t.get('AIQ', 0) - t.get('GOI', 0):
            print('  Req. Of != Cap - AIQ - GOI'); sp += 1
        print(f"[{'OK' if sp == 0 else 'CHECK'}] {s['name']}: {len(s['colleges'])} colleges, {sp} problem(s)")
        problems += sp
    return problems


# ----------------------------------------------------------------------------
# 3. EXCEL OUTPUT
# ----------------------------------------------------------------------------
TITLE = ["GOVERNMENT OF MAHARASHTRA", "State Common Entrance Test Cell,Maharashtra, Mumbai",
         "Admissions to Health Science Courses, NEET(UG)-2026-2027",
         "Provisional Seat Matrix for MBBS/BDS Courses in Govt./Govt. Aided/Corp./Private/Minority Medical/Dental Colleges – Round 3",
         "Printed On :06/10/2026"]
LEGEND = ["Legends: Min.:Minority Seats, I.Q:Institute Quota EW:EWS HA: Hilly Area MK:MKB, VJA:VJ(A), NTB:NT(B)/NT1, NTC:NT(C)/NT2, NTD:NT(D)/NT3, ORP: Orphan (Institutional-IWC/IOG), ORPC: Orphan (Outside Institution-NIN)",
          "Note:",
          "1. Seat position is subject to change as per information from Institute/ MUHS, Nashik/NMC/DCI/Order of Hon’ble Courts/ State/ Central Government.",
          "2. Government / CET Cell Commissioner may change this seat position before allotment of seats in CAP round(s).",
          "3. Seat will be allotted after receiving affiliation from MUHS, Nashik or respective Council."]
thin, hair, dbl = Side(style='thin'), Side(style='hair', color='808080'), Side(style='double')
GREY = PatternFill('solid', fgColor='D9D9D9')
Fnt = lambda **k: Font(name='Arial', size=k.pop('size', 9), **k)


def kind_of(name):
    if 'UNIVERSITY' in name:
        return 'univ'
    return 'govt' if name.startswith('GOVT') else 'pvt'


def sheet_name(name):
    course = name.split(':')[1].strip()
    return {'govt': 'Govt', 'pvt': 'Pvt', 'univ': 'Pvt Univ'}[kind_of(name)] + ' ' + course


def columns(kind):
    c = [('Sr', 'Sr', '', 5), ('Code', 'Code', '', 7), ('name', 'College', '', 58), ('gov', 'Govt/', 'Aided', 7), ('Cap', 'Int.', 'Cap.', 7)]
    if kind == 'govt':
        c.append(('AIQ', 'AIQ', '', 6))
    c += [('min', 'Min', '', 9), ('GOI', 'GOI', '', 5), ('label', '', '', 10)]
    for k in ['SC', 'ST', 'VJA', 'NTB', 'NTC', 'NTD', 'OBC']:
        c.append((k, '', k, 5))
    c += [('CTot', '', 'Tot.', 5), ('SEB', '', 'SEB', 5), ('EWS', '', 'EWS', 5), ('OPEN', '', 'OPEN', 6),
          ('StateTot', 'State', 'Tot.', 6), ('StateTotCol', '', '', 6),
          ('D1', '', 'D1', 5), ('D2', '', 'D2', 5), ('D3', '', 'D3', 5), ('MK', '', 'MK', 5)]
    c += [('IQ', '', 'NRI', 5), ('AIQ2', '', 'AIQ', 5)] if kind == 'univ' else [('IQ', '', 'IQ', 5)]
    c += [('OthTot', '', 'Tot.', 6), ('OthTotCol', '', '', 6), ('Bal', 'Tot.', 'Bal.', 6), ('BalCol', '', '', 6)]
    return c


def write_data_sheet(wb, sec):
    kind, nm = kind_of(sec['name']), sheet_name(sec['name'])
    ws = wb.create_sheet(nm)
    cols = columns(kind)
    ix = {k: i + 1 for i, (k, *_) in enumerate(cols)}
    n = len(cols)
    for i, t in enumerate(TITLE):
        ws.cell(1 + i, 1, t).font = Fnt(bold=True, size=11 if i < 3 else 9)
    ws.cell(7, ix['label'], sec['name']).font = Fnt(bold=True, size=10)
    ws.cell(7, ix['D1'], 'Status : Original').font = Fnt(bold=True)
    for k, h1, h2, w in cols:
        for r, h in ((8, h1), (9, h2)):
            c = ws.cell(r, ix[k], h or None)
            c.font = Fnt(bold=True); c.fill = GREY; c.alignment = Alignment(horizontal='center')
        ws.column_dimensions[L(ix[k])].width = w
    ws.cell(8, ix['SC'], '<- - - - - - - - - - -State - - - - - - ->').alignment = Alignment(horizontal='left')
    ws.cell(8, ix['D1'], '<- - -Other Resevation - ->').alignment = Alignment(horizontal='left')
    for c in range(1, n + 1):
        ws.cell(8, c).fill = ws.cell(9, c).fill = GREY
        ws.cell(8, c).border = Border(top=thin); ws.cell(9, c).border = Border(bottom=thin)
    for k, txt in (('label', 'Category'),
                   ('StateTotCol', 'College total of State Tot. (printed once per college, first row)'),
                   ('OthTotCol', 'College total of Other Reservation Tot. (printed once per college, first row)'),
                   ('BalCol', 'College total of Bal. (printed once per college, first row)')):
        ws.cell(9, ix[k]).comment = Comment(txt, 'Converter')

    def put(d, r, bold=False):
        for k, v in d.items():
            if v is None or v == '':
                continue
            c = ws.cell(r, ix[k], v)
            c.font = Fnt(bold=bold)
            c.alignment = Alignment(horizontal='left' if k in ('name', 'label', 'min') else 'center')

    r, blocks = 10, []
    for col in sec['colleges']:
        r0 = r
        for i, row in enumerate(col['rows']):
            d = {'label': row['label'], **row['vals']}
            if i == 0:
                d.update(Sr=col['sr'], Code=col['code'], name=col['name'], gov=col['gov'], min=col['min'] or None)
            put(d, r); r += 1
        for c in range(1, n + 1):
            ws.cell(r - 1, c).border = Border(bottom=hair)
        blocks.append((r0, r - 1))
    d1, t0 = r - 1, r
    for i, row in enumerate(sec['total']['rows']):
        d = {'label': row['label'], **row['vals']}
        if i == 0:
            d['name'] = '* * Total * *'
        put(d, r, bold=True); r += 1
    for c in range(1, n + 1):
        ws.cell(t0, c).border = Border(top=dbl); ws.cell(r - 1, c).border = Border(bottom=dbl)
    reqrow = r
    ws.cell(reqrow, ix['name'], sec['req_text']).font = Fnt(bold=True)
    ws.cell(reqrow, ix['name']).alignment = Alignment(horizontal='right')
    keys = ['SC', 'ST', 'VJA', 'NTB', 'NTC', 'NTD', 'OBC', 'CTot', 'SEB', 'EWS', 'OPEN', 'StateTot', 'D1', 'D2', 'D3', 'MK', 'IQ'] + (['AIQ2'] if kind == 'univ' else []) + ['OthTot']
    labs = ['SC', 'ST', 'VJA', 'NTB', 'NTC', 'NTD', 'OBC', 'Tot.', 'SEB', 'EWS', 'OPEN', 'Tot.', 'D1', 'D2', 'D3', 'MK'] + (['NRI', 'AIQ'] if kind == 'univ' else ['IQ']) + ['Tot.']
    for k, t in zip(keys, labs):
        c = ws.cell(reqrow, ix[k], t); c.font = Fnt(bold=True); c.alignment = Alignment(horizontal='center')
    r = reqrow + 2
    for t in LEGEND:
        ws.cell(r, 1, t).font = Fnt(size=8, bold=(t == 'Note:')); r += 1
    for rr in range(8, reqrow):                       # the PDF's "|" dividers
        for k in ('StateTot', 'D1', 'Bal'):
            c = ws.cell(rr, ix[k]); b = c.border
            c.border = Border(left=thin, top=b.top, bottom=b.bottom, right=b.right)
    ws.freeze_panes = ws.cell(10, ix['label'])
    ws.page_setup.orientation = 'landscape'; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return dict(name=nm, ix=ix, d0=10, d1=d1, t0=t0, labels=[x['label'] for x in sec['total']['rows']],
                reqrow=reqrow, blocks=blocks, kind=kind)


def write_checks_sheet(wb, metas):
    ck = wb.create_sheet('Checks', 0)
    ck.column_dimensions['A'].width = 34; ck.column_dimensions['B'].width = 48
    for i in range(3, 40):
        ck.column_dimensions[L(i)].width = 11
    ck['A1'] = 'Validation of PDF → Excel conversion (live formulas on the data sheets; every difference should be 0)'
    ck['A1'].font = Fnt(bold=True, size=12)
    hdr = ['Section / sheet', 'Colleges counted', 'Row: Tot. ≠ SC..OBC', 'Row: State Tot ≠ Tot+SEB+EWS+OPEN',
           'Row: Other Tot ≠ D1..IQ (NRI+AIQ)', 'Row: Bal ≠ State Tot + Other Tot', 'College totals mismatches',
           'Printed Total-row mismatches', 'Int.Cap/AIQ/GOI/Req mismatches', 'Status']
    for i, h in enumerate(hdr):
        c = ck.cell(3, 1 + i, h); c.font = Fnt(bold=True); c.fill = GREY
        c.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center')
    ck.row_dimensions[3].height = 52
    q = lambda n: "'" + n + "'"
    cur = 4 + len(metas) + 3
    refs = {}
    for m in metas:
        n, ix, d0, d1 = m['name'], m['ix'], m['d0'], m['d1']
        col = lambda k: L(ix[k])
        rng = lambda k: f"{q(n)}!${col(k)}${d0}:${col(k)}${d1}"
        ck.cell(cur, 1, f'{n} – college-level checks').font = Fnt(bold=True, size=10)
        for i, t in enumerate(['Code', 'College', 'State Tot (college) printed', 'Σ rows', 'Diff', 'Other Tot (college) printed',
                               'Σ rows', 'Diff', 'Bal (college) printed', 'Σ rows', 'Diff', 'Bal(col) − (State col + Other col)']):
            c = ck.cell(cur + 1, 1 + i, t); c.font = Fnt(bold=True); c.fill = GREY
            c.alignment = Alignment(wrap_text=True, horizontal='center')
        ck.row_dimensions[cur + 1].height = 40
        r = cur + 2; cs = r
        for a, b in m['blocks']:
            P = lambda k: f"{q(n)}!{col(k)}{a}"
            Rg = lambda k: f"{q(n)}!{col(k)}{a}:{col(k)}{b}"
            ck.cell(r, 1, f"={P('Code')}"); ck.cell(r, 2, f"={P('name')}")
            ck.cell(r, 3, f"=N({P('StateTotCol')})"); ck.cell(r, 4, f"=SUM({Rg('StateTot')})"); ck.cell(r, 5, f"=C{r}-D{r}")
            ck.cell(r, 6, f"=N({P('OthTotCol')})"); ck.cell(r, 7, f"=SUM({Rg('OthTot')})"); ck.cell(r, 8, f"=F{r}-G{r}")
            ck.cell(r, 9, f"=N({P('BalCol')})"); ck.cell(r, 10, f"=SUM({Rg('Bal')})"); ck.cell(r, 11, f"=I{r}-J{r}")
            ck.cell(r, 12, f"=I{r}-(C{r}+F{r})")
            for c in range(1, 13):
                ck.cell(r, c).font = Fnt()
            r += 1
        ce = r - 1
        ck.cell(r, 2, 'Mismatching colleges →').font = Fnt(bold=True)
        ck.cell(r, 5, '=' + '+'.join(f'SUMPRODUCT(--({c}{cs}:{c}{ce}<>0))' for c in 'EHKL')).font = Fnt(bold=True)
        colmis = f'E{r}'; r += 2

        ck.cell(r, 1, f'{n} – printed Total row vs Σ of college rows (Σ − printed)').font = Fnt(bold=True, size=10); r += 1
        mk = ['SC', 'ST', 'VJA', 'NTB', 'NTC', 'NTD', 'OBC', 'CTot', 'SEB', 'EWS', 'OPEN', 'StateTot', 'D1', 'D2', 'D3', 'MK', 'IQ'] + (['AIQ2'] if m['kind'] == 'univ' else []) + ['OthTot', 'Bal']
        mh = ['SC', 'ST', 'VJA', 'NTB', 'NTC', 'NTD', 'OBC', 'Tot.', 'SEB', 'EWS', 'OPEN', 'State Tot.', 'D1', 'D2', 'D3', 'MK'] + (['NRI', 'AIQ'] if m['kind'] == 'univ' else ['IQ']) + ['Other Tot.', 'Bal.']
        ck.cell(r, 1, 'Category').font = Fnt(bold=True)
        for i, t in enumerate(mh):
            c = ck.cell(r, 2 + i, t); c.font = Fnt(bold=True); c.fill = GREY; c.alignment = Alignment(horizontal='center')
        r += 1; ms = r
        for j, lab in enumerate(m['labels']):
            ck.cell(r, 1, lab).font = Fnt()
            for i, k in enumerate(mk):
                ck.cell(r, 2 + i, f"=SUMIF({rng('label')},$A{r},{rng(k)})-N({q(n)}!{col(k)}{m['t0'] + j})").font = Fnt()
            r += 1
        me = r - 1
        ck.cell(r, 1, 'Mismatching cells →').font = Fnt(bold=True)
        ck.cell(r, 2, f"=SUMPRODUCT(--(B{ms}:{L(1 + len(mk))}{me}<>0))").font = Fnt(bold=True)
        totmis = f'B{r}'; r += 1

        ck.cell(r, 1, 'Σ Int. Cap − printed total'); ck.cell(r, 2, f"=SUM({rng('Cap')})-N({q(n)}!{col('Cap')}{m['t0']})"); capc = f'B{r}'; r += 1
        aiqc = None
        if m['kind'] == 'govt':
            ck.cell(r, 1, 'Σ AIQ − printed total'); ck.cell(r, 2, f"=SUM({rng('AIQ')})-N({q(n)}!{col('AIQ')}{m['t0']})"); aiqc = f'B{r}'; r += 1
        ck.cell(r, 1, 'Σ GOI − printed total'); ck.cell(r, 2, f"=SUM({rng('GOI')})-N({q(n)}!{col('GOI')}{m['t0']})"); goic = f'B{r}'; r += 1
        ck.cell(r, 1, 'Req. Of − (Cap − AIQ − GOI)')
        rc = f"{q(n)}!{col('name')}{m['reqrow']}"
        aiqterm = f"-N({q(n)}!{col('AIQ')}{m['t0']})" if m['kind'] == 'govt' else ''
        ck.cell(r, 2, f'=VALUE(TRIM(MID({rc},FIND(":",{rc})+1,20)))-(N({q(n)}!{col("Cap")}{m["t0"]}){aiqterm}-N({q(n)}!{col("GOI")}{m["t0"]}))')
        reqc = f'B{r}'; r += 1
        capmis = '+'.join(f'SUMPRODUCT(--({x}<>0))' for x in (capc, goic, reqc) + ((aiqc,) if aiqc else ()))
        refs[n] = (colmis, totmis, capmis)
        cur = r + 2

    for i, m in enumerate(metas):
        n, ix, d0, d1 = m['name'], m['ix'], m['d0'], m['d1']
        col = lambda k: L(ix[k])
        rng = lambda k: f"{q(n)}!${col(k)}${d0}:${col(k)}${d1}"
        r = 4 + i
        colmis, totmis, capmis = refs[n]
        oth = ['D1', 'D2', 'D3', 'MK', 'IQ'] + (['AIQ2'] if m['kind'] == 'univ' else [])
        ck.cell(r, 1, n)
        ck.cell(r, 2, f"=COUNT({rng('Sr')})")
        ck.cell(r, 3, f"=SUMPRODUCT(--({rng('CTot')}<>{'+'.join(rng(k) for k in ['SC', 'ST', 'VJA', 'NTB', 'NTC', 'NTD', 'OBC'])}))")
        ck.cell(r, 4, f"=SUMPRODUCT(--({rng('StateTot')}<>{rng('CTot')}+{rng('SEB')}+{rng('EWS')}+{rng('OPEN')}))")
        ck.cell(r, 5, f"=SUMPRODUCT(--({rng('OthTot')}<>{'+'.join(rng(k) for k in oth)}))")
        ck.cell(r, 6, f"=SUMPRODUCT(--({rng('Bal')}<>{rng('StateTot')}+{rng('OthTot')}))")
        ck.cell(r, 7, f'={colmis}'); ck.cell(r, 8, f'={totmis}'); ck.cell(r, 9, f'={capmis}')
        ck.cell(r, 10, f'=IF(SUM(C{r}:I{r})=0,"OK","CHECK")')
        for c in range(1, 11):
            ck.cell(r, c).font = Fnt(bold=c in (1, 10))
            ck.cell(r, c).alignment = Alignment(horizontal='center' if c > 1 else 'left')
    rg = f'J4:J{3 + len(metas)}'
    ck.conditional_formatting.add(rg, CellIsRule(operator='equal', formula=['"OK"'], fill=PatternFill('solid', bgColor='C6EFCE', fgColor='C6EFCE')))
    ck.conditional_formatting.add(rg, CellIsRule(operator='equal', formula=['"CHECK"'], fill=PatternFill('solid', bgColor='FFC7CE', fgColor='FFC7CE')))
    ck.freeze_panes = 'B4'


def build_workbook(sections, out_path):
    wb = Workbook(); wb.remove(wb.active)
    metas = [write_data_sheet(wb, s) for s in sections]
    write_checks_sheet(wb, metas)
    wb.save(out_path)



# ----------------------------------------------------------------------------
# 4. "LONG" OUTPUT: one row per (college, row-type, category) seat entry
# ----------------------------------------------------------------------------
LONG_CATS = ['SC', 'ST', 'VJA', 'NTB', 'NTC', 'NTD', 'OBC', 'SEB', 'EWS', 'OPEN', 'D1', 'D2', 'D3', 'MK', 'IQ', 'AIQ2']
LABEL_ORDER = ['GEN.', 'WOM.', 'PWD.', 'ORP.', 'ORPC', 'GEN HA', 'WOMEN HA']


def long_rows(sec, keep_zeros=False):
    """Rows for one section, grouped category by category (SC block, ST block, ...).
    GOI is a college-level number, so it is written ONCE per college (first row of that college)."""
    kind = kind_of(sec['name'])
    cat_name = {'IQ': 'NRI' if kind == 'univ' else 'IQ', 'AIQ2': 'AIQ'}
    cat_name.update(sec.get('catnames', {}))             # use the PDF's own labels (VJ/N1.. or VJA/NTB..)
    goi_done = set()
    out = []
    for cat in LONG_CATS:
        for col in sec['colleges']:
            rows = sorted(col['rows'], key=lambda r: LABEL_ORDER.index(r['label']))
            for r in rows:
                v = r['vals'].get(cat)
                if v is None or (v == 0 and not keep_zeros):
                    continue
                goi = None
                if col['code'] not in goi_done:
                    goi = next((x['vals'].get('GOI') for x in col['rows'] if x['vals'].get('GOI') is not None), None)
                    goi_done.add(col['code'])
                out.append([col['sr'], col['code'], col['name'], col['gov'], col['min'] or None,
                            goi, r['label'], cat_name.get(cat, cat), v, sheet_name(sec['name'])])
    return out


def build_long_workbook(sections, out_path, keep_zeros=False, section_col=False, csv=False):
    wb = Workbook(); wb.remove(wb.active)
    courses = collections.OrderedDict()
    for s in sections:                                   # MBBS sections together, BDS sections together
        courses.setdefault(s['name'].split(':')[1].strip(), []).append(s)
    header = ['Sr', 'Code', 'College', 'Aided', 'Min', 'GOI', '', 'Category', 'Seats'] + (['Section'] if section_col else [])
    widths = [6, 7, 62, 7, 9, 6, 10, 10, 7, 16]
    lines = []
    for course, secs in courses.items():
        ws = wb.create_sheet(course)
        ws.append(header)
        for sec in secs:
            for row in long_rows(sec, keep_zeros):
                ws.append(row if section_col else row[:-1])
                lines.append('\t'.join('' if x is None else str(x) for x in (row if section_col else row[:-1])))
        for i in range(len(header)):
            ws.column_dimensions[L(i + 1)].width = widths[i]
            ws.cell(1, i + 1).font = Fnt(bold=True); ws.cell(1, i + 1).fill = GREY
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font = Fnt()
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = f'A1:{L(len(header))}{ws.max_row}'
        print(f'  {course}: {ws.max_row - 1} rows')
    wb.save(out_path)
    if csv:
        path = re.sub(r'\.xlsx$', '', out_path, flags=re.I) + '.txt'
        with open(path, 'w', encoding='utf-8') as f:
            f.write('\t'.join(header) + '\n' + '\n'.join(lines) + '\n')
        print('  also wrote', path)


# ----------------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser(description='Convert Maharashtra CET seat-matrix PDF to Excel')
    ap.add_argument('pdf'); ap.add_argument('xlsx', nargs='?')
    ap.add_argument('--debug-grid', action='store_true', help='print x-positions of numbers (to tune GRID)')
    ap.add_argument('--long', action='store_true', help='one row per seat entry (Sr, Code, College, Aided, Min, GOI, type, Category, Seats)')
    ap.add_argument('--keep-zeros', action='store_true', help='with --long: keep entries printed as 0')
    ap.add_argument('--section-col', action='store_true', help='with --long: add a Section column')
    ap.add_argument('--csv', action='store_true', help='with --long: also write a tab-separated text file')
    a = ap.parse_args()
    out = a.xlsx or re.sub(r'\.pdf$', '', a.pdf, flags=re.I) + '.xlsx'
    sections = parse_pdf(a.pdf, a.debug_grid)
    sections = [x for x in sections if x['colleges'] and x['total']]
    if not sections:
        print('ERROR: no Maharashtra CET seat-matrix tables (GOVT./PVT. COLLEGES: MBBS/BDS) were found in this PDF.\n'
              '       This script only reads the "Provisional Seat Matrix ... State CET Cell, Maharashtra" report.\n'
              '       Other PDFs (for example a Punjab "SEATS LIST ... Vacant Seats" list) need a different converter.')
        return 1
    problems = validate(sections)
    if a.long:
        build_long_workbook(sections, out, a.keep_zeros, a.section_col, a.csv)
    else:
        build_workbook(sections, out)
    print(f'\nSaved {out}   ({"all checks passed" if problems == 0 else str(problems) + " problem(s) - see above"})')
    if not a.long:
        print("Open the file in Excel/LibreOffice; the 'Checks' sheet recalculates on open "
              "(or run: soffice --headless --convert-to xlsx --outdir fixed <file>).")


if __name__ == '__main__':
    sys.exit(main())
